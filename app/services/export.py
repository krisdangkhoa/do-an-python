"""Xuat bao cao thu chi ra tep Excel."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.models import Transaction, TransactionType, User
from app.services import report as svc

HEADER_FILL = PatternFill("solid", fgColor="003A70")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = "#,##0"


def _header(ws, row, titles):
    for i, title in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def period_label(year, month) -> str:
    if not year:
        return "Toàn bộ"
    return f"Tháng {month}/{year}" if month else f"Năm {year}"


def build_workbook(db: Session, user: User, year=None, month=None) -> bytes:
    """Tao tep Excel, chi lay du lieu cua nguoi dung dang dang nhap."""
    start, end = svc.period_bounds(year, month)
    s = svc.summary(db, user.id, year, month)
    wb = Workbook()

    # --- Trang 1: Tong quan ---
    ws = wb.active
    ws.title = "Tổng quan"
    ws["A1"] = "BÁO CÁO THU - CHI"
    ws["A1"].font = Font(bold=True, size=14, color="003A70")
    ws["A2"] = f"Người dùng: {user.full_name}"
    ws["A3"] = f"Kỳ báo cáo: {period_label(year, month)}"
    _header(ws, 5, ["Chỉ số", "Giá trị"])
    rows = [("Tổng thu nhập", s["income"]), ("Tổng chi tiêu", s["expense"]),
            ("Số dư", s["balance"]), ("Số giao dịch", s["count"])]
    for r, (label, value) in enumerate(rows, 6):
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=value)
        if label != "Số giao dịch":
            c.number_format = MONEY
    _widths(ws, [22, 20])

    # --- Trang 2: Danh sach giao dich ---
    ws2 = wb.create_sheet("Giao dịch")
    _header(ws2, 1, ["Ngày", "Loại", "Danh mục", "Số tiền", "Ghi chú"])
    q = (
        db.query(Transaction)
        .options(joinedload(Transaction.category))
        .filter(Transaction.user_id == user.id)
    )
    if start:
        q = q.filter(Transaction.date >= start)
    if end:
        q = q.filter(Transaction.date <= end)
    for r, t in enumerate(q.order_by(Transaction.date).all(), 2):
        ws2.cell(row=r, column=1, value=t.date).number_format = "dd/mm/yyyy hh:mm"
        ws2.cell(row=r, column=2, value="Thu nhập" if t.type == TransactionType.INCOME.value else "Chi tiêu")
        ws2.cell(row=r, column=3, value=t.category.name)
        ws2.cell(row=r, column=4, value=t.amount).number_format = MONEY
        ws2.cell(row=r, column=5, value=t.note or "")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    _widths(ws2, [18, 12, 18, 16, 36])

    # --- Trang 3: Ty trong theo danh muc ---
    ws3 = wb.create_sheet("Theo danh mục")
    _header(ws3, 1, ["Loại", "Danh mục", "Số tiền", "Tỷ trọng (%)"])
    r = 2
    for type_, label in ((TransactionType.EXPENSE.value, "Chi tiêu"),
                         (TransactionType.INCOME.value, "Thu nhập")):
        for row in svc.by_category(db, user.id, type_, year, month):
            ws3.cell(row=r, column=1, value=label)
            ws3.cell(row=r, column=2, value=row["name"])
            ws3.cell(row=r, column=3, value=row["amount"]).number_format = MONEY
            ws3.cell(row=r, column=4, value=row["percent"])
            r += 1
    _widths(ws3, [12, 20, 16, 14])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()