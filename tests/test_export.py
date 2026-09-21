"""Kiem thu chuc nang xuat Excel."""
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.models import Category, User
from tests.conftest import login, register

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _add(client, db, username, amount, note):
    cat = (
        db.query(Category).join(User)
        .filter(User.username == username, Category.name == "Ăn uống")
        .first()
    )
    client.post("/transactions", data={
        "amount": str(amount), "type": "expense",
        "date": datetime.now().strftime("%Y-%m-%dT%H:%M"),
        "note": note, "category_id": str(cat.id),
    }, follow_redirects=False)


def test_xuat_excel_dung_dinh_dang_va_so_lieu(auth_client, db_session):
    """XE-01: Tep Excel dung dinh dang, dung so lieu."""
    _add(auth_client, db_session, "user1", 250000, "An trua")
    now = datetime.now()

    r = auth_client.get(f"/report/export?year={now.year}&month={now.month}")

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(XLSX)
    assert "attachment" in r.headers["content-disposition"]

    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Tổng quan", "Giao dịch", "Theo danh mục"]
    assert wb["Giao dịch"].max_row == 2
    assert wb["Giao dịch"]["D2"].value == 250000
    assert wb["Tổng quan"]["B7"].value == 250000  # tong chi tieu


def test_xuat_excel_khong_lan_du_lieu_nguoi_khac(client, db_session):
    """XE-02: Tep cua A khong chua giao dich cua B."""
    for name, amount in (("userA", 100000), ("userB", 200000)):
        register(client, name)
        login(client, name)
        _add(client, db_session, name, amount, f"Cua {name}")

    login(client, "userA")
    ws = load_workbook(BytesIO(client.get("/report/export").content))["Giao dịch"]
    notes = [ws.cell(row=i, column=5).value for i in range(2, ws.max_row + 1)]

    assert notes == ["Cua userA"]


def test_chua_dang_nhap_khong_xuat_duoc(client):
    """XE-03: Phai dang nhap moi tai duoc."""
    r = client.get("/report/export", follow_redirects=False)
    assert r.status_code == 303